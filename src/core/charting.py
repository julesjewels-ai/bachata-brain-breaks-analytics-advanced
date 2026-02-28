"""
Charting module for Excel reports.
Encapsulates chart configuration and building logic using OpenPyXL.
"""
from typing import Optional
from pydantic import BaseModel, Field, ConfigDict
from openpyxl.chart import BarChart, Reference
from openpyxl.worksheet.worksheet import Worksheet


class ChartDataLocation(BaseModel):
    """Defines the location of data for the chart."""
    min_col: int
    min_row: int
    max_col: int
    max_row: int
    title_from_data: bool = True

    # Optional explicit location for categories (labels)
    cats_min_col: Optional[int] = None
    cats_min_row: Optional[int] = None
    cats_max_row: Optional[int] = None


class ChartConfig(BaseModel):
    """Configuration for chart styling and dimensions."""
    title: str = Field(..., description="Title of the chart")
    x_axis_title: str = Field(..., description="Label for X-Axis")
    y_axis_title: str = Field(..., description="Label for Y-Axis")
    width: float = Field(default=15.0, description="Chart width in cm")
    height: float = Field(default=10.0, description="Chart height in cm")
    style: int = Field(default=10, description="Excel chart style index")

    model_config = ConfigDict(arbitrary_types_allowed=True)


class ChartBuilder:
    """Builds OpenPyXL charts based on configuration."""

    def __init__(self, worksheet: Worksheet):
        self.ws = worksheet

    def add_bar_chart(self,
                      data_loc: ChartDataLocation,
                      config: ChartConfig,
                      anchor: str = "E2") -> None:
        """
        Creates a bar chart and adds it to the worksheet.

        Args:
            data_loc: Location of the data for the chart.
            config: Chart styling configuration.
            anchor: The cell where the top-left corner of the chart will be placed.
        """
        chart = BarChart()
        chart.title = config.title
        chart.style = config.style
        chart.y_axis.title = config.y_axis_title
        chart.x_axis.title = config.x_axis_title

        # OpenPyXL expects generic numbers or strings for dimensions
        # using # type: ignore to suppress mypy errors as per memory
        chart.width = config.width  # type: ignore
        chart.height = config.height  # type: ignore

        # Data References
        data = Reference(
            self.ws,
            min_col=data_loc.min_col,
            min_row=data_loc.min_row,
            max_col=data_loc.max_col,
            max_row=data_loc.max_row
        )

        chart.add_data(data, titles_from_data=data_loc.title_from_data)

        # Categories (Labels)
        # Use explicit location if provided, otherwise default to column left
        # of data
        cats_min_col = data_loc.cats_min_col if data_loc.cats_min_col is not None else data_loc.min_col - 1

        # Default rows usually match data rows (adjusted for header if needed)
        default_min_row = data_loc.min_row + \
            1 if data_loc.title_from_data else data_loc.min_row
        cats_min_row = data_loc.cats_min_row if data_loc.cats_min_row is not None else default_min_row

        cats_max_row = data_loc.cats_max_row if data_loc.cats_max_row is not None else data_loc.max_row

        cats = Reference(
            self.ws,
            min_col=cats_min_col,
            min_row=cats_min_row,
            max_row=cats_max_row
        )
        chart.set_categories(cats)

        self.ws.add_chart(chart, anchor)
