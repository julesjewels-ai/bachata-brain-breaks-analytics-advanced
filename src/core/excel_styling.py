"""
Excel styling utilities.
Handles low-level styling, formatting, and layout for Excel reports.
"""
from typing import Dict
from numbers import Number
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import DataBarRule
from openpyxl.worksheet.worksheet import Worksheet


class ExcelStyler:
    """Helper class for applying standardized styles to Excel worksheets."""

    HEADER_FONT = Font(bold=True, color="FFFFFF")
    HEADER_FILL = PatternFill(start_color="4F81BD", fill_type="solid")

    @staticmethod
    def _get_formatted_number_length(val, fmt) -> int:
        """Helper to calculate length of formatted number."""
        # Thousands separator (e.g., #,##0)
        if '#,##0' in fmt:
            precision = 2 if '.00' in fmt else 0
            return len(f"{val:,.{precision}f}")

        # Percentage (e.g., 0.00%)
        if '0.00%' in fmt or '0.00"%"' in fmt:
            return len(f"{val:.2f}%")

        return len(str(val))

    @staticmethod
    def estimate_cell_width(cell) -> int:
        """
        Estimates the display width of a cell based on value and number format.
        """
        if cell.value is None:
            return 0

        val = cell.value
        fmt = cell.number_format

        # Return early if not a number with a format
        if not (isinstance(val, Number) and fmt):
            return len(str(val))

        return ExcelStyler._get_formatted_number_length(val, fmt)

    @staticmethod
    def adjust_column_widths(ws: Worksheet):
        """Auto-adjusts column widths based on content length with min/max constraints."""
        min_width = 10
        max_width = 50
        for col in ws.columns:
            # Calculate max length of data in column
            max_length = 0
            for cell in col:
                cell_width = ExcelStyler.estimate_cell_width(cell)
                max_length = max(max_length, cell_width)

            # Apply padding and clamp between min and max
            adjusted_width = max(min_width, min(max_length + 2, max_width))
            if col[0].column:
                ws.column_dimensions[get_column_letter(col[0].column)].width = adjusted_width

    @staticmethod
    def get_header_map(ws: Worksheet) -> Dict[str, int]:
        """Returns a map of header name to column index (1-based)."""
        return {str(cell.value): cell.column for cell in ws[1] if cell.value is not None}

    @staticmethod
    def apply_header_style(ws: Worksheet):
        """Applies standard header styling (Bold, Centered, Blue) and freezes panes."""
        for cell in ws[1]:
            cell.font = ExcelStyler.HEADER_FONT
            cell.fill = ExcelStyler.HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = 'A2'

    @staticmethod
    def apply_number_formats(ws: Worksheet):
        """Applies number formatting to specific columns."""
        # Map column headers to their respective formats
        format_map = {
            'Views': '#,##0',
            'Retention (%)': '0.00"%"'
        }

        headers = ExcelStyler.get_header_map(ws)

        for header, fmt in format_map.items():
            if header in headers:
                col_idx = headers[header]
                # Apply format to all cells in the column (skipping header)
                for row in range(2, ws.max_row + 1):
                    ws.cell(row=row, column=col_idx).number_format = fmt

    @staticmethod
    def apply_conditional_formatting(ws: Worksheet):
        """Applies data bars to visualization columns."""
        # Define rules
        # Blue for Views, Green for Retention
        rules = {
            'Views': DataBarRule(start_type='min', end_type='max', color="638EC6"),
            'Retention (%)': DataBarRule(start_type='min', end_type='max', color="63C384")
        }

        headers = ExcelStyler.get_header_map(ws)

        for header, rule in rules.items():
            if header in headers:
                col_letter = get_column_letter(headers[header])
                # Apply to the entire column data range (e.g. C2:C100)
                # Ensure we have data
                if ws.max_row > 1:
                    range_ref = f"{col_letter}2:{col_letter}{ws.max_row}"
                    ws.conditional_formatting.add(range_ref, rule)
