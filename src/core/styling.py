"""
Styling module for Excel reports.
Handles all formatting, color, and layout logic for Excel generation.
"""
from typing import Dict
from numbers import Number
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import DataBarRule
from openpyxl.worksheet.worksheet import Worksheet

class ExcelStyler:
    """
    Handles styling and formatting logic for Excel reports.
    Follows Single Responsibility Principle by isolating presentation logic.
    """

    # Constants for consistent styling
    HEADER_FONT = Font(bold=True, color="FFFFFF")
    HEADER_FILL = PatternFill(start_color="4F81BD", fill_type="solid")

    COLOR_VIEWS_BAR = "638EC6"
    COLOR_RETENTION_BAR = "63C384"

    COLUMN_WIDTH_MIN = 10
    COLUMN_WIDTH_MAX = 50

    @staticmethod
    def _get_header_map(ws: Worksheet) -> Dict[str, int]:
        """Returns a map of header name to column index (1-based)."""
        return {str(cell.value): cell.column for cell in ws[1] if cell.value is not None}

    def apply_header_style(self, ws: Worksheet) -> None:
        """Applies standard header styling (Bold, Centered, Blue) and freezes panes."""
        for cell in ws[1]:
            cell.font = self.HEADER_FONT
            cell.fill = self.HEADER_FILL
            cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.freeze_panes = 'A2'

    def estimate_cell_width(self, cell) -> float:
        """
        Estimates the display width of a cell based on value and number format.
        Applies a multiplier for bold text (headers).
        """
        if cell.value is None:
            return 0.0

        val = cell.value
        number_format = cell.number_format

        width = 0.0

        # Handle formatted numbers
        if isinstance(val, Number) and number_format:
            # Thousands separator (e.g., #,##0)
            if '#,##0' in number_format:
                precision = 2 if '.00' in number_format else 0
                width = len(f"{val:,.{precision}f}")

            # Percentage (e.g., 0.00%)
            elif '0.00"%"' in number_format or '0.00%' in number_format:
                width = len(f"{val:.2f}%")
            else:
                width = len(str(val))
        else:
            width = len(str(val))

        # Apply multiplier for bold text (e.g., headers)
        if cell.font and cell.font.bold:
            width *= 1.2

        return width

    def adjust_column_widths(self, ws: Worksheet) -> None:
        """Auto-adjusts column widths based on content length with min/max constraints."""
        for col in ws.columns:
            # Calculate max length of data in column
            max_length = 0.0
            for cell in col:
                cell_width = self.estimate_cell_width(cell)
                max_length = max(max_length, cell_width)

            # Apply padding and clamp between min and max
            adjusted_width = max(self.COLUMN_WIDTH_MIN, min(max_length + 2, self.COLUMN_WIDTH_MAX))
            ws.column_dimensions[get_column_letter(col[0].column)].width = adjusted_width

    def apply_number_formats(self, ws: Worksheet) -> None:
        """Applies number formatting to specific columns."""
        # Map column headers to their respective formats
        format_map = {
            'Views': '#,##0',
            'Retention (%)': '0.00"%"'
        }

        headers = self._get_header_map(ws)

        for header, fmt in format_map.items():
            if header in headers:
                col_idx = headers[header]
                # Apply format to all cells in the column (skipping header)
                for row in range(2, ws.max_row + 1):
                    ws.cell(row=row, column=col_idx).number_format = fmt

    def apply_conditional_formatting(self, ws: Worksheet) -> None:
        """Applies data bars to visualization columns."""
        rules = {
            'Views': DataBarRule(start_type='min', end_type='max', color=self.COLOR_VIEWS_BAR),
            'Retention (%)': DataBarRule(start_type='min', end_type='max', color=self.COLOR_RETENTION_BAR)
        }

        headers = self._get_header_map(ws)

        for header, rule in rules.items():
            if header in headers:
                col_letter = get_column_letter(headers[header])
                # Ensure we have data
                if ws.max_row > 1:
                    range_ref = f"{col_letter}2:{col_letter}{ws.max_row}"
                    ws.conditional_formatting.add(range_ref, rule)

    def format_strategy_sheet(self, ws: Worksheet) -> None:
        """Applies specific formatting for the Strategy sheet."""
        self.apply_header_style(ws)
        ws.column_dimensions['A'].width = 100
        # Ensure the cell exists before accessing properties
        # This assumes data is already written to A2
        if ws.max_row >= 2:
            ws['A2'].alignment = Alignment(wrap_text=True, horizontal='left', vertical='top')
