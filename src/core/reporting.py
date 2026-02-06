"""
Reporting module for generating Excel reports.
Handles styling and formatting logic for Excel output.
"""
from typing import Dict
import logging
import pandas as pd
from openpyxl.styles import Font, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.drawing.image import Image as XLImage
from PIL import Image as PILImage
from pydantic import BaseModel, Field, ValidationError, field_validator
import re

from src.core.charting import ChartBuilder, ChartConfig, ChartDataLocation
from src.core.visualization import MatplotlibVisualizer
from src.core.excel_styling import ExcelStyler

logger = logging.getLogger(__name__)

class ReportConfig(BaseModel):
    """Configuration for report generation validation."""
    filepath: str = Field(..., description="Path to save the Excel report")

    @field_validator('filepath')
    @classmethod
    def validate_filepath(cls, v: str) -> str:
        if not v.endswith('.xlsx'):
            raise ValueError("File must be an Excel (.xlsx) file")
        if '..' in v:
            raise ValueError("Path traversal detected")
        if not re.match(r'^[\w\-. /]+$', v):
            raise ValueError("File path contains invalid characters")
        return v


class ExcelReportGenerator:
    """Generates styled Excel reports for analytics data."""

    # Mapping from DataFrame columns to Excel headers
    COLUMN_MAPPING = {
        'video_id': 'Video ID',
        'title': 'Video Title',
        'views': 'Views',
        'retention_avg_pct': 'Retention (%)',
        'type': 'Type',
        'publish_date': 'Publish Date'
    }

    def _create_anomaly_sheet(self, writer, v_type: str, df: pd.DataFrame) -> Worksheet:
        """Creates and styles a sheet for anomalies."""
        sheet_name = f"{v_type} Anomalies"
        display_df = df.rename(columns=self.COLUMN_MAPPING)
        display_df.to_excel(writer, sheet_name=sheet_name, index=False)
        ws = writer.sheets[sheet_name]
        ExcelStyler.apply_header_style(ws)
        ExcelStyler.apply_number_formats(ws)
        ExcelStyler.apply_conditional_formatting(ws)
        ExcelStyler.adjust_column_widths(ws)
        return ws

    def _add_anomaly_chart(self, ws: Worksheet, v_type: str) -> None:
        """Adds a bar chart to the anomaly sheet."""
        headers = ExcelStyler.get_header_map(ws)

        if 'Views' not in headers or 'Video Title' not in headers:
            return

        # Only add chart if there is data
        if ws.max_row <= 1:
            return

        views_col = headers['Views']
        title_col = headers['Video Title']
        max_row = ws.max_row
        max_col = ws.max_column

        chart_builder = ChartBuilder(ws)
        data_loc = ChartDataLocation(
            min_col=views_col,
            min_row=1,  # Include header for series name
            max_col=views_col,
            max_row=max_row,
            title_from_data=True,
            cats_min_col=title_col
        )
        chart_config = ChartConfig(
            title=f"Top {v_type} Views",
            x_axis_title="Video Title",
            y_axis_title="Views"
        )

        # Dynamic anchor: 2 columns to the right of the table
        anchor_col = get_column_letter(max_col + 2)

        chart_builder.add_bar_chart(
            data_loc=data_loc,
            config=chart_config,
            anchor=f"{anchor_col}2"
        )

    def _add_strategy_sheet(self, writer, strategy: str) -> None:
        """Creates the strategy analysis sheet."""
        pd.DataFrame({'Gemini Analysis': [strategy]}).to_excel(
            writer, sheet_name="Strategy", index=False
        )
        ws_strat = writer.sheets["Strategy"]
        ExcelStyler.apply_header_style(ws_strat)
        ws_strat.column_dimensions['A'].width = 100
        ws_strat['A2'].alignment = Alignment(wrap_text=True, horizontal='left', vertical='top')

    def _add_visual_insights(self, writer, anomalies: Dict[str, pd.DataFrame]) -> None:
        """Generates and embeds visual insights chart."""
        # Combine all anomalies to one DF for visualization
        all_anomalies = pd.concat(anomalies.values()) if anomalies else pd.DataFrame()
        if not all_anomalies.empty and 'views' in all_anomalies.columns and 'retention_avg_pct' in all_anomalies.columns:
            visualizer = MatplotlibVisualizer()
            try:
                img_stream = visualizer.generate_chart(
                    all_anomalies,
                    title="Views vs Retention Correlation",
                    x_col="retention_avg_pct",
                    y_col="views"
                )

                # Create sheet
                ws_viz = writer.book.create_sheet("Visual Insights")

                # Embed Image
                # OpenPyXL Image requires a path or PIL Image object
                pil_img = PILImage.open(img_stream)
                img = XLImage(pil_img)
                ws_viz.add_image(img, "A1")

                # Add description
                ws_viz["A25"] = "Scatter plot showing relationship between Audience Retention and View Count."
                ws_viz["A25"].font = Font(italic=True, color="555555")

            except Exception as e:
                # Log or handle error without crashing report
                logger.warning(f"Failed to generate visualization: {e}")

    def generate_excel(self,
                       anomalies: Dict[str, pd.DataFrame],
                       strategy: str,
                       filepath: str):
        """Creates an Excel report with anomalies and strategy analysis."""
        try:
            config = ReportConfig(filepath=filepath)
            safe_path = config.filepath
        except ValidationError as e:
            raise ValueError(f"Security validation failed: {e}")

        with pd.ExcelWriter(safe_path, engine='openpyxl') as writer:
            # 1. Anomalies Sheets
            for v_type, df in anomalies.items():
                if not df.empty:
                    ws = self._create_anomaly_sheet(writer, v_type, df)
                    self._add_anomaly_chart(ws, v_type)

            # 2. Strategy Sheet
            self._add_strategy_sheet(writer, strategy)

            # 3. Visual Insights (Embedded Matplotlib)
            self._add_visual_insights(writer, anomalies)
