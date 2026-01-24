import os
import pandas as pd
import pytest
from openpyxl import load_workbook
from src.core.reporting import ExcelReportGenerator
from src.core.visualization import MatplotlibVisualizer
from src.core.services import AnalyticsReportService

@pytest.fixture
def mock_anomalies():
    return {
        'Long': pd.DataFrame({
            'video_id': ['1', '2'],
            'title': ['Video A', 'Video B'],
            'views': [1000, 2000],
            'retention_avg_pct': [50.5, 60.0],
            'type': ['Long', 'Long'],
            'publish_date': ['2023-01-01', '2023-01-02']
        })
    }

def test_full_report_generation(mock_anomalies, tmp_path):
    # Setup
    filepath = str(tmp_path / "test_report.xlsx")
    viz = MatplotlibVisualizer()
    gen = ExcelReportGenerator()
    service = AnalyticsReportService(gen, viz)

    # Execute
    service.create_comprehensive_report(mock_anomalies, "Test Strategy", filepath)

    # Verify File Exists
    assert os.path.exists(filepath)

    # Verify Excel Content
    wb = load_workbook(filepath)
    assert "Long Anomalies" in wb.sheetnames
    ws = wb["Long Anomalies"]

    # Verify Data
    assert ws['A1'].value == "Video ID"
    assert ws['C2'].value == 1000

    # Verify Image Exists
    # In openpyxl, images are stored in the _images list of the worksheet
    assert len(ws._images) > 0, "Expected at least one image (heatmap) to be embedded"
