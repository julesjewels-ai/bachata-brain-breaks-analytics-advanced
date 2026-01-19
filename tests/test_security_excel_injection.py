"""
Security tests for Excel Injection vulnerabilities.
Verifies that user inputs starting with dangerous characters (=, @, +, -) are properly sanitized.
"""
import pytest
import pandas as pd
import openpyxl
from src.core.reporting import ExcelReportGenerator, ExcelSanitizer

def test_excel_injection_sanitization(tmp_path):
    """
    Test that malicious inputs (formulas) are sanitized by prefixing with a single quote.
    """
    malicious_inputs = ["=1+1", "@SUM(1,1)", "+10-5", "-20"]
    safe_inputs = ["Normal Text", "123", "User=Name"]

    df = pd.DataFrame({
        'title': malicious_inputs + safe_inputs,
        'views': [100] * (len(malicious_inputs) + len(safe_inputs)),
        'retention_avg_pct': [50.0] * (len(malicious_inputs) + len(safe_inputs))
    })

    anomalies = {'SecurityTest': df}
    strategy = "=DANGEROUS_STRATEGY()"
    output_file = tmp_path / "security_test.xlsx"

    generator = ExcelReportGenerator()
    generator.generate_excel(anomalies, strategy, str(output_file))

    # Verify Output
    wb = openpyxl.load_workbook(output_file)

    # Check Anomalies Sheet
    ws = wb['SecurityTest Anomalies']

    # Check malicious inputs
    for i, inp in enumerate(malicious_inputs):
        cell_val = ws[f'A{i+2}'].value
        expected = f"'{inp}"
        assert cell_val == expected, f"Malicious input '{inp}' was not sanitized! Got: {cell_val}"

    # Check safe inputs
    start_safe = len(malicious_inputs)
    for i, inp in enumerate(safe_inputs):
        cell_val = ws[f'A{start_safe + i + 2}'].value
        assert cell_val == inp, f"Safe input '{inp}' was altered! Got: {cell_val}"

    # Check Strategy Sheet
    ws_strat = wb['Strategy']
    strat_val = ws_strat['A2'].value
    assert strat_val == f"'{strategy}", "Strategy input was not sanitized!"

def test_sanitizer_logic():
    """Unit test for ExcelSanitizer class."""
    assert ExcelSanitizer.sanitize("=SUM(1,1)") == "'=SUM(1,1)"
    assert ExcelSanitizer.sanitize("@cmd") == "'@cmd"
    assert ExcelSanitizer.sanitize("+1") == "'+1"
    assert ExcelSanitizer.sanitize("-1") == "'-1"
    assert ExcelSanitizer.sanitize("Normal") == "Normal"
    assert ExcelSanitizer.sanitize(123) == 123
    assert ExcelSanitizer.sanitize(None) is None

if __name__ == "__main__":
    pytest.main([__file__])
