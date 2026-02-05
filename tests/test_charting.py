"""
Tests for the charting module.
"""
from unittest.mock import MagicMock
from src.core.charting import ChartBuilder, ChartConfig, ChartDataLocation
from openpyxl.chart import BarChart


def test_chart_builder_add_bar_chart():
    """Test that add_bar_chart configures and adds a chart to the worksheet."""
    mock_ws = MagicMock()
    builder = ChartBuilder(mock_ws)

    data_loc = ChartDataLocation(
        min_col=2,
        min_row=1,
        max_col=2,
        max_row=5,
        title_from_data=True,
        cats_min_col=1  # Explicit category column
    )

    config = ChartConfig(
        title="Test Chart",
        x_axis_title="X Axis",
        y_axis_title="Y Axis"
    )

    builder.add_bar_chart(data_loc, config)

    # Verify add_chart was called
    assert mock_ws.add_chart.called

    # Get the chart object passed to add_chart
    chart_arg = mock_ws.add_chart.call_args[0][0]
    assert isinstance(chart_arg, BarChart)

    # OpenPyXL chart titles are objects; access text content via nested
    # properties
    # chart.title (Title) -> tx (Text) -> rich (RichText) -> p (Paragraphs) ->
    # r (Run) -> t (Text)
    try:
        title_text = chart_arg.title.tx.rich.p[0].r[0].t
    except AttributeError:
        # Fallback if structure is different (e.g. simple string assignment
        # simulation)
        title_text = str(chart_arg.title)

    assert title_text == "Test Chart"

    # Check X Axis Title
    try:
        x_title = chart_arg.x_axis.title.tx.rich.p[0].r[0].t
    except AttributeError:
        x_title = str(chart_arg.x_axis.title)
    assert x_title == "X Axis"

    # Check Y Axis Title
    try:
        y_title = chart_arg.y_axis.title.tx.rich.p[0].r[0].t
    except AttributeError:
        y_title = str(chart_arg.y_axis.title)
    assert y_title == "Y Axis"

    # Verify anchor
    anchor_arg = mock_ws.add_chart.call_args[0][1]
    assert anchor_arg == "E2"
