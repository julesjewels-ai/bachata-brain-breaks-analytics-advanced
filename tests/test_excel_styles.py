import pytest
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.formatting.rule import DataBarRule
from src.core.excel_styles import ExcelStyler

@pytest.fixture
def ws():
    wb = Workbook()
    ws = wb.active
    return ws

def test_estimate_cell_width_string(ws):
    ws['A1'] = "Hello World"
    width = ExcelStyler.estimate_cell_width(ws['A1'])
    assert width == 11

def test_estimate_cell_width_number_format(ws):
    ws['A1'] = 12345.67
    ws['A1'].number_format = '#,##0.00'
    # 12,345.67 -> 9 chars
    width = ExcelStyler.estimate_cell_width(ws['A1'])
    assert width == 9

def test_adjust_column_widths(ws):
    ws['A1'] = "Header"
    ws['A2'] = "Very Long Content In This Cell"

    ExcelStyler.adjust_column_widths(ws, min_width=10, max_width=50)

    col_width = ws.column_dimensions['A'].width
    assert col_width > 20
    assert col_width <= 50

def test_apply_header_style(ws):
    ws['A1'] = "Header 1"
    ws['B1'] = "Header 2"

    ExcelStyler.apply_header_style(ws)

    # Check Font
    assert ws['A1'].font.bold is True
    # OpenPyXL might return color as ThemeColor or RGB.
    # If RGB, it is usually "00FFFFFF" (ARGB) or just "FFFFFF".
    # Let's check simply if it's set.
    assert ws['A1'].font.color is not None

    # Check Fill
    assert ws['A1'].fill.patternType == 'solid'

    # Check Freeze Panes
    assert ws.freeze_panes == 'A2'

def test_apply_number_formats(ws):
    ws['A1'] = "Value"
    ws['A2'] = 1000

    format_map = {'Value': '#,##0'}
    ExcelStyler.apply_number_formats(ws, format_map)

    assert ws['A2'].number_format == '#,##0'

def test_apply_conditional_formatting(ws):
    ws['A1'] = "Value"
    ws['A2'] = 10
    ws['A3'] = 20

    rule = DataBarRule(start_type='min', end_type='max', color="638EC6")
    rules = {'Value': rule}

    ExcelStyler.apply_conditional_formatting(ws, rules)

    # Check if rule exists
    count = 0
    for cf in ws.conditional_formatting:
        for r in cf.rules:
            if r.type == 'dataBar':
                count += 1
    assert count >= 1
